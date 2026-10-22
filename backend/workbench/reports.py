from io import BytesIO
import html
import json
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn

FIELDS = [("id", "Issue ID"), ("rule_id", "Rule ID"), ("severity", "Severity"), ("category", "Clause category"), ("observed_language", "Observed language"), ("source_anchor", "Source anchor"), ("explanation", "Explanation"), ("preferred_position", "Preferred contract position"), ("negotiating_response", "Suggested negotiating response"), ("status", "Reviewer status"), ("notes", "Reviewer notes"), ("confidence", "Confidence or uncertainty"), ("origin", "Analysis layer")]

def value(issue, field):
    if field == "observed_language" and issue["absence"]:
        return "ABSENCE: No matching provision located. No quotation is asserted."
    if field == "source_anchor":
        a = issue[field]
        return f"{a['id']}; {a['location']}; extracted offsets [{a['start']}, {a['end']})" if a else "ABSENCE: No source passage"
    return str(issue[field] or "(none)")

def export(review, document, format):
    report = {"report_version": 1, "author": "Ethan", "notice": "Contract review aid, not legal advice. Attorney review is required.", "document": {k: document[k] for k in ["id", "name", "sha256", "format"]}, **review}
    if format == "json":
        return json.dumps(report, ensure_ascii=False, indent=2).encode(), "application/json", "json"
    title = "Contract review report"
    intro = f"Document: {document['name']}\nPlaybook: {review['playbook_name']}\nAI status: {review['ai_status']}\n{report['notice']}"
    if format == "markdown":
        lines = [f"# {title}", "", html.escape(intro), "", f"Source SHA256: {document['sha256']}", ""]
        for issue in review["issues"]:
            lines.extend([f"## {html.escape(issue['category'])}", ""])
            for field, label in FIELDS:
                lines.extend([f"**{label}:** {html.escape(value(issue, field))}", ""])
        lines += ["## Rule checks", "", *[f"- {c['rule_id']}: {c['outcome']} {html.escape(json.dumps(c['properties']))}" for c in review["checks"]]]
        return "\n".join(lines).encode(), "text/markdown; charset=utf-8", "md"
    doc = Document()
    doc.core_properties.author = "Ethan"
    doc.core_properties.last_modified_by = "Ethan"
    doc.core_properties.title = title
    doc.core_properties.subject = "Source-linked commercial contract review"
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    for style in doc.styles:
        if style.type == 1:
            style.font.color.rgb = RGBColor(0, 0, 0)
            for border in list(style.element.iter(qn("w:pBdr"))):
                border.getparent().remove(border)
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Heading 1"].font.size = Pt(13)
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.05
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.75)
    doc.add_heading(title, 0)
    doc.add_paragraph(intro)
    doc.add_paragraph(f"Source SHA256: {document['sha256']}")
    for issue in review["issues"]:
        doc.add_heading(issue["category"], 1)
        for field, label in FIELDS:
            p = doc.add_paragraph()
            p.add_run(label + ": ").bold = True
            p.add_run(value(issue, field))
    doc.add_heading("Rule checks", 1)
    for check in review["checks"]:
        doc.add_paragraph(f"{check['rule_id']}: {check['outcome']} {json.dumps(check['properties'])}")
    out = BytesIO()
    doc.save(out)
    return out.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
