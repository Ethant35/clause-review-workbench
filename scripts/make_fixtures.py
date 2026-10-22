"""Generate real synthetic DOCX/PDF files for API and browser acceptance."""
from pathlib import Path
from docx import Document
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts/fixtures"

def make():
    OUT.mkdir(parents=True, exist_ok=True)
    source = (ROOT / "corpus/contracts/04-renewal-and-unilateral-audit.txt").read_text()
    doc = Document()
    doc.core_properties.author = "Ethan"
    doc.core_properties.last_modified_by = "Ethan"
    doc.core_properties.title = "Synthetic commercial agreement"
    doc.add_heading("Synthetic commercial agreement", 0)
    for text in source.split("\n\n"):
        if text.strip(): doc.add_paragraph(text.strip())
    doc.save(OUT / "synthetic.docx")
    pdf = pymupdf.open()
    # One clause per text block, ample space, two real pages.
    for group in [source.split("\n\n")[:9], source.split("\n\n")[9:]]:
        page = pdf.new_page(width=612, height=792)
        y=55
        for text in group:
            if text.strip():
                # Standard PDF font uses ASCII apostrophe; extraction expected against this text.
                text=text.strip().replace("’", "'")
                remaining = page.insert_textbox(pymupdf.Rect(54,y,558,y+65), text, fontsize=10, fontname="helv")
                assert remaining >= 0, "Fixture paragraph clipped"
                y+=76
    pdf.set_metadata({"author":"Ethan", "title":"Synthetic commercial agreement"})
    pdf.save(OUT / "synthetic.pdf")
    scan=pymupdf.open()
    scan.new_page().draw_rect(pymupdf.Rect(50,50,300,300), color=(0,0,0), fill=(0.5,0.5,0.5))
    scan.save(OUT / "image-only.pdf")
    return OUT
if __name__ == "__main__": print(make())
