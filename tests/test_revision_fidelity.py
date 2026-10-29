"""Source fidelity, revision isolation, and comparison regressions. Author: Ethan."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
import pymupdf
import pytest

from backend.workbench import store
from backend.workbench.comparison import compare_reviews
from backend.workbench.main import app, save_playbook
from backend.workbench.parsing import ParseError, extract
from scripts.evaluate_cuad import overlap_size

ROOT = Path(__file__).parents[1]
CAP = "Each party's aggregate liability will not exceed the fees paid in the preceding 12 months."
EXCEPTION = "Indemnification obligations are excluded from the foregoing cap."


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "DATA", tmp_path)
    monkeypatch.setattr(store, "RETENTION_DAYS", 0)
    store.initialize()
    save_playbook("commercial", (ROOT / "playbooks/commercial.yaml").read_text())
    return TestClient(app, base_url="http://127.0.0.1")


def uploaded(client, text, previous=None):
    result = client.post("/api/documents", params={"revision_of": previous} if previous else {},
                         files={"file": ("contract.txt", text.encode(), "text/plain")})
    assert result.status_code == 201, result.text
    return result.json()


def reviewed(client, document):
    result = client.post(f"/api/documents/{document['id']}/reviews", json={"playbook_id": "commercial"})
    assert result.status_code == 201, result.text
    return result.json()


def comparison(client, old, new):
    result = client.get(f"/api/reviews/{new['id']}/comparison", params={"baseline_review_id": old["id"]})
    assert result.status_code == 200, result.text
    return result.json()


def test_exception_revision_preserves_originals_reviews_and_decisions(client):
    first = uploaded(client, CAP + "\n" + EXCEPTION)
    old = reviewed(client, first)
    cap = next(i for i in old["issues"] if i["rule_id"] == "liability-cap")
    patch = client.patch(f"/api/reviews/{old['id']}/issues/{cap['id']}", json={"status": "rejected", "notes": "Earlier cap scope decision"})
    assert patch.status_code == 200
    frozen = store.get("reviews", old["id"])
    new_doc = uploaded(client, CAP + "\n" + EXCEPTION.replace("Indemnification", "Confidentiality"), first["id"])
    assert new_doc["previous_document_id"] == first["id"]
    new = reviewed(client, new_doc)
    assert all(i["status"] == "unresolved" and i["notes"] == "" for i in new["issues"])
    delta = comparison(client, old, new)
    row = next(r for r in delta["rules"] if r["rule_id"] == "liability-cap")
    assert row["status"] == "CHANGED" and row["revisit_prior_decisions"]
    assert row["prior_decisions"][0]["notes"] == "Earlier cap scope decision"
    assert delta["text_changes"][0]["before"][0]["text"] == EXCEPTION
    assert store.get("reviews", old["id"]) == frozen
    assert client.get(f"/api/documents/{first['id']}/original").content.decode() == CAP + "\n" + EXCEPTION
    assert client.get(f"/api/documents/{new_doc['id']}/original").content.decode() == CAP + "\n" + EXCEPTION.replace("Indemnification", "Confidentiality")
    export = client.get(f"/api/reviews/{new['id']}/comparison", params={"baseline_review_id": old["id"], "download": True})
    assert export.json() == delta and "attachment" in export.headers["content-disposition"]
    assert client.get(f"/api/reviews/{old['id']}/export?format=json").json()["issues"] == client.get(f"/api/reviews/{old['id']}").json()["issues"]


def test_deleted_exception_is_no_longer_flagged_not_legal_clearance(client):
    a = reviewed(client, uploaded(client, CAP + "\n" + EXCEPTION))
    b = reviewed(client, uploaded(client, CAP))
    delta = comparison(client, a, b)
    assert next(r for r in delta["rules"] if r["rule_id"] == "liability-cap")["status"] == "NO_LONGER_FLAGGED"
    assert any("not a legal clearance" in w for w in delta["warnings"])


def test_narrow_fraud_exception_still_requires_revisiting_the_prior_decision(client):
    old = reviewed(client, uploaded(client, CAP + "\n" + EXCEPTION))
    issue = next(i for i in old["issues"] if i["rule_id"] == "liability-cap")
    client.patch(f"/api/reviews/{old['id']}/issues/{issue['id']}", json={"status": "rejected"})
    new = reviewed(client, uploaded(client, CAP + "\n" + EXCEPTION.replace("Indemnification", "Fraud")))
    result = comparison(client, old, new)
    row = next(r for r in result["rules"] if r["rule_id"] == "liability-cap")
    assert row["status"] == "NO_LONGER_FLAGGED"
    assert row["revisit_prior_decisions"]
    assert any("not a legal clearance" in warning for warning in result["warnings"])


def test_separate_limit_change_is_visible_even_when_cap_words_do_not_change(client):
    text = CAP + "\n" + EXCEPTION + "\nIndemnification liability is capped at $200,000."
    a = reviewed(client, uploaded(client, text))
    b = reviewed(client, uploaded(client, text.replace("$200,000", "$900,000")))
    delta = comparison(client, a, b)
    row = next(r for r in delta["rules"] if r["rule_id"] == "liability-cap")
    assert row["status"] == "CHANGED"
    assert any("$900,000" in anchor["text"] for issue in row["after_issues"] for relation in issue.get("interactions", []) for anchor in relation["evidence"])


def test_analysis_revision_changes_are_not_contract_edits(client):
    doc = uploaded(client, CAP)
    a, b = reviewed(client, doc), reviewed(client, doc)
    historical = store.get("reviews", a["id"])
    historical["review_revision"] = 1
    with store.db() as c:
        c.execute("UPDATE reviews SET data=? WHERE id=?", (json.dumps(historical), a["id"]))
    delta = comparison(client, a, b)
    assert not delta["same_checks"]
    assert all(r["status"] == "NOT_COMPARABLE" for r in delta["rules"])


def test_identical_text_different_original_bytes_requires_revisit(client):
    doc = uploaded(client, EXCEPTION)
    a = reviewed(client, doc)
    issue = a["issues"][0]
    client.patch(f"/api/reviews/{a['id']}/issues/{issue['id']}", json={"status": "accepted"})
    b = reviewed(client, uploaded(client, "\n" + EXCEPTION + "\n"))
    delta = comparison(client, a, b)
    assert delta["same_extracted_text"] and not delta["same_original"]
    assert not delta["text_changes"]
    assert next(r for r in delta["rules"] if r["rule_id"] == issue["rule_id"])["revisit_prior_decisions"]
    assert any("original files differ" in w for w in delta["warnings"])


def test_unchanged_anchor_does_not_hide_unmodeled_dependency_edit(client):
    first = uploaded(client, EXCEPTION + "\nSpecial Schedule controls all terms.")
    a = reviewed(client, first)
    issue = a["issues"][0]
    client.patch(f"/api/reviews/{a['id']}/issues/{issue['id']}", json={"notes": "Check schedule"})
    b = reviewed(client, uploaded(client, EXCEPTION + "\nSpecial Schedule no longer controls terms."))
    row = next(r for r in comparison(client, a, b)["rules"] if r["rule_id"] == issue["rule_id"])
    assert row["revisit_prior_decisions"]


def test_policy_changes_are_not_reported_as_contract_effects(client):
    doc = uploaded(client, CAP)
    a = reviewed(client, doc)
    policy = store.get("playbooks", "commercial")["yaml"]
    save_playbook("commercial", policy.replace("severity: high", "severity: low"))
    b = reviewed(client, doc)
    delta = comparison(client, a, b)
    assert not delta["same_policy"]
    assert all(r["status"] == "NOT_COMPARABLE" for r in delta["rules"])


def test_same_bytes_new_review_never_copies_decisions(client):
    doc = uploaded(client, EXCEPTION)
    a = reviewed(client, doc)
    client.patch(f"/api/reviews/{a['id']}/issues/{a['issues'][0]['id']}", json={"status": "accepted"})
    b = reviewed(client, doc)
    assert all(i["status"] == "unresolved" for i in b["issues"])
    delta = comparison(client, a, b)
    assert all(r["status"] == "UNCHANGED" and not r["revisit_prior_decisions"] for r in delta["rules"])


@pytest.mark.parametrize("corruption", ["original", "extraction"])
def test_comparison_rejects_changed_source(client, corruption):
    doc = uploaded(client, CAP)
    a, b = reviewed(client, doc), reviewed(client, doc)
    if corruption == "original":
        store.source_path(doc["id"]).write_bytes(b"replaced")
    else:
        doc["text"] += "unrecorded change"
        with store.db() as c:
            c.execute("UPDATE documents SET data=? WHERE id=?", (json.dumps(doc), doc["id"]))
    assert client.get(f"/api/reviews/{b['id']}/comparison", params={"baseline_review_id": a["id"]}).status_code == 409


def test_deleted_parent_does_not_cascade_to_revision(client):
    doc = uploaded(client, CAP)
    old = reviewed(client, doc)
    revision = uploaded(client, EXCEPTION, doc["id"])
    current = reviewed(client, revision)
    assert client.delete(f"/api/documents/{doc['id']}").status_code == 200
    assert client.get(f"/api/reviews/{current['id']}").status_code == 200
    assert client.get(f"/api/documents/{revision['id']}/original").status_code == 200
    assert client.get(f"/api/reviews/{current['id']}/comparison", params={"baseline_review_id": old["id"]}).status_code == 404
    assert client.post("/api/documents?revision_of=missing", files={"file": ("x.txt", b"Terms")}).status_code == 404


def test_repeated_moved_text_and_numbers_remain_in_delta(client):
    doc = uploaded(client, "Section 1: Fees are $10.\nRepeated clause.\nRepeated clause.\nSection 2: Cap is $20.")
    old = reviewed(client, doc)
    revised = uploaded(client, "Section 2: Cap is $21.\nRepeated clause.\nSection 1: Fees are $10.\nRepeated clause.")
    new = reviewed(client, revised)
    delta = comparison(client, old, new)
    assert any("$20" in b["text"] for group in delta["text_changes"] for b in group["before"])
    assert any("$21" in b["text"] for group in delta["text_changes"] for b in group["after"])


def docx_bytes(doc):
    stream = BytesIO()
    doc.core_properties.author = doc.core_properties.last_modified_by = "Ethan"
    doc.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize("where", ["direct", "style"])
def test_automatic_numbering_cannot_silently_disappear(where):
    doc = Document()
    p = doc.add_paragraph("Indemnity is subject to Section 2.")
    if where == "style":
        p.style = "List Number"
    else:
        props = p._p.get_or_add_pPr()
        numbering = OxmlElement("w:numPr")
        identifier = OxmlElement("w:numId")
        identifier.set(qn("w:val"), "1")
        numbering.append(identifier)
        props.append(numbering)
    with pytest.raises(ParseError, match="automatic numbering"):
        extract(docx_bytes(doc))


@pytest.mark.parametrize("where", ["header", "footer"])
def test_terms_outside_docx_body_cannot_silently_disappear(where):
    doc = Document()
    doc.add_paragraph(CAP)
    getattr(doc.sections[0], where).paragraphs[0].text = EXCEPTION
    with pytest.raises(ParseError, match="headers or footers"):
        extract(docx_bytes(doc))


def test_literal_numbering_and_table_positions_are_exact():
    doc = Document()
    doc.add_paragraph("2. " + CAP)
    table = doc.add_table(rows=2, cols=2)
    for row, texts in zip(table.rows, [("Category", "Limit"), ("Indemnity", "$200")]):
        for cell, text in zip(row.cells, texts):
            cell.text = text
    result = extract(docx_bytes(doc))
    cell = next(b for b in result["blocks"] if b["text"] == "$200")
    assert (cell["table"], cell["row"], cell["cell"]) == (1, 2, 2)
    assert result["blocks"][0]["text"] == "2. " + CAP
    assert all(result["text"][b["start"]:b["end"]] == b["text"] for b in result["blocks"])


@pytest.mark.parametrize("complexity", ["columns", "annotation", "simple"])
def test_pdf_layout_coordinates_and_visual_gate(client, complexity):
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((40, 60), "First column terms.")
    if complexity == "columns":
        page.insert_text((330, 60), "Second column exceptions.")
    elif complexity == "annotation":
        page.add_text_annot((100, 100), "Indemnity is uncapped.")
    data = doc.tobytes()
    doc.close()
    result = extract(data)
    assert all(len(b["bbox_points"]) == 4 for b in result["blocks"])
    assert result["visual_check"]["status"] == ("NOT_REQUIRED" if complexity == "simple" else "NEEDS_VISUAL_CHECK")
    uploaded_doc = client.post("/api/documents", files={"file": ("layout.pdf", data)}).json()
    review = client.post(f"/api/documents/{uploaded_doc['id']}/reviews", json={"playbook_id": "commercial"})
    assert review.status_code == (201 if complexity == "simple" else 409)


def test_external_label_overlap_does_not_double_count_evidence():
    assert overlap_size((10, 20), [(0, 15), (12, 18), (18, 25)]) == 10
    assert overlap_size((10, 20), [(20, 30), (0, 10)]) == 0
    assert overlap_size((10, 20), [(12, 14), (16, 18)]) == 4
