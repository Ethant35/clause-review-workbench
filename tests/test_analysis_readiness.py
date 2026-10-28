"""Independent regression and metamorphic checks for extraction and analysis."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import unittest
from zipfile import ZipFile

import yaml
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pymupdf
from hypothesis import given, settings, strategies as st

from workbench.engine import analyze
from workbench.models import Playbook
from workbench.parsing import ParseError, extract
from workbench.safety import EXTRACTION_REVISION, REVIEW_REVISION, review_view


ROOT = Path(__file__).resolve().parents[1]
PLAYBOOK = Playbook.model_validate(
    yaml.safe_load((ROOT / "playbooks/commercial.yaml").read_text())
)
CAP_ONLY = PLAYBOOK.model_copy(
    update={"rules": [rule for rule in PLAYBOOK.rules if rule.id == "liability-cap"]}
)
NDA_PLAYBOOK = Playbook.model_validate(
    yaml.safe_load((ROOT / "playbooks/nda.yaml").read_text())
)


def cap_review(blocks):
    document = extract(("\n\n".join(blocks)).encode())
    issues, checks = analyze(document, CAP_ONLY)
    linked = [
        relation
        for issue in issues
        for relation in issue.get("interactions", [])
        if relation["kind"] == "cap-exception"
    ]
    return document, checks[0], linked


class CapScopeRegression(unittest.TestCase):
    def test_exclusion_of_foregoing_cap_in_another_paragraph_requires_review(self):
        document, check, relations = cap_review(
            [
                "Each party's aggregate liability will not exceed the Fees paid in the preceding twelve months.",
                "Indemnification obligations are excluded from the foregoing cap.",
            ]
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertEqual(1, len(relations))
        self.assertEqual(
            {document["blocks"][0]["id"], document["blocks"][1]["id"]},
            {anchor["id"] for anchor in relations[0]["evidence"]},
        )
        self.assertTrue(
            all(
                document["text"][anchor["start"] : anchor["end"]] == anchor["text"]
                for anchor in relations[0]["evidence"]
            )
        )

    @settings(max_examples=80, deadline=None)
    @given(st.lists(st.integers(min_value=0, max_value=99), max_size=6))
    def test_unrelated_paragraph_insertions_preserve_cap_finding_and_evidence(self, fillers):
        cap = "Each party's aggregate liability will not exceed the Fees paid in the preceding twelve months."
        exception = "Indemnification obligations are not subject to the liability cap."
        baseline = [cap, exception]
        blocks = [f"Neutral recital {number}." for number in fillers]
        insert_at = len(fillers) // 2
        expanded = [*blocks[:insert_at], cap, *blocks[insert_at:], exception]

        _, baseline_check, baseline_relations = cap_review(baseline)
        document, expanded_check, expanded_relations = cap_review(expanded)

        self.assertEqual(baseline_check["outcome"], expanded_check["outcome"])
        self.assertEqual("NEEDS_REVIEW", expanded_check["outcome"])
        self.assertEqual(1, len(baseline_relations))
        self.assertEqual(1, len(expanded_relations))
        self.assertEqual(
            {cap, exception},
            {anchor["text"] for anchor in expanded_relations[0]["evidence"]},
        )
        self.assertTrue(
            all(
                document["text"][anchor["start"] : anchor["end"]] == anchor["text"]
                for anchor in expanded_relations[0]["evidence"]
            )
        )
    def test_non_exclusion_in_another_paragraph_is_not_a_cap_exception(self):
        _, check, relations = cap_review(
            [
                "Each party's aggregate liability will not exceed the Fees paid in the preceding twelve months.",
                "The foregoing cap does not exclude indemnification obligations.",
            ]
        )

        self.assertEqual("COMPLIANT", check["outcome"])
        self.assertEqual([], relations)

    def test_foregoing_limitations_carveout_in_another_paragraph_requires_review(self):
        _, check, relations = cap_review(
            [
                "Each party's aggregate liability will not exceed the Fees paid in the preceding twelve months.",
                "The foregoing limitations do not apply to either party's indemnification obligations.",
            ]
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertEqual(1, len(relations))

    def test_not_subject_to_cap_in_another_paragraph_requires_review(self):
        _, check, relations = cap_review(
            [
                "Each party's aggregate liability will not exceed the Fees paid in the preceding twelve months.",
                "Indemnification obligations are not subject to the liability cap.",
            ]
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertEqual(1, len(relations))

    def test_foregoing_cap_will_not_apply_in_another_paragraph_requires_review(self):
        document, check, relations = cap_review(
            [
                "Each party's aggregate liability will not exceed the Fees paid in the preceding twelve months.",
                "The foregoing cap will not apply to indemnification obligations.",
            ]
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertEqual(1, len(relations))
        self.assertEqual("indemnification", relations[0]["exception_scope"])
        self.assertEqual(
            {block["id"] for block in document["blocks"]},
            {anchor["id"] for anchor in relations[0]["evidence"]},
        )

    def test_same_block_liability_cap_will_not_apply_requires_review(self):
        document = extract(
            b"Aggregate liability will not exceed Fees paid in the preceding twelve months, "
            b"and the liability cap will not apply to indemnification obligations."
        )
        issues, checks = analyze(document, CAP_ONLY)
        issue = next(issue for issue in issues if issue["rule_id"] == "liability-cap")

        self.assertEqual("NEEDS_REVIEW", checks[0]["outcome"])
        relation = next(item for item in issue["interactions"] if item["kind"] == "cap-exception")
        self.assertEqual("indemnification obligations", relation["exception_scope"])
        self.assertTrue(all(document["text"][anchor["start"]:anchor["end"]] == anchor["text"]
                            for anchor in relation["evidence"]))

    def test_will_not_apply_fraud_exception_keeps_configured_narrow_exception(self):
        _, check, relations = cap_review(
            [
                "Each party's aggregate liability will not exceed the Fees paid in the preceding twelve months.",
                "The foregoing cap will not apply to fraud or willful misconduct.",
            ]
        )

        self.assertEqual("COMPLIANT", check["outcome"])
        self.assertEqual([], relations)

    def test_will_not_exclude_is_not_a_cap_exception(self):
        _, check, relations = cap_review(
            [
                "Each party's aggregate liability will not exceed the Fees paid in the preceding twelve months.",
                "The foregoing cap will not exclude indemnification obligations.",
            ]
        )

        self.assertEqual("COMPLIANT", check["outcome"])
        self.assertEqual([], relations)

    def test_configured_narrow_exception_stays_compliant(self):
        _, check, relations = cap_review(
            [
                "Each party's aggregate liability will not exceed the Fees paid in the preceding twelve months.",
                "The liability cap does not apply to fraud or willful misconduct.",
            ]
        )

        self.assertEqual("COMPLIANT", check["outcome"])
        self.assertEqual([], relations)

    def test_exception_list_with_indemnity_scope_requires_review(self):
        document, check, relations = cap_review(
            [
                "Each party's aggregate liability will not exceed the Fees paid in the preceding twelve months.",
                "The liability cap does not apply to fraud, willful misconduct, or indemnification obligations.",
            ]
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertEqual(1, len(relations))
        self.assertEqual(
            {block["text"] for block in document["blocks"]},
            {anchor["text"] for anchor in relations[0]["evidence"]},
        )

    def test_same_block_negated_carveout_cannot_confirm_fee_cap(self):
        document = extract(
            b"Aggregate liability will not exceed Fees paid in the preceding twelve months, "
            b"and indemnification obligations are not excluded from this cap."
        )
        issues, checks = analyze(document, CAP_ONLY)

        self.assertEqual("NEEDS_REVIEW", checks[0]["outcome"])
        self.assertFalse(checks[0]["properties"]["explicit_fee_cap"])
        cap_issue = next(issue for issue in issues if issue["rule_id"] == "liability-cap")
        self.assertEqual(document["blocks"][0], cap_issue["source_anchor"])

    def test_named_exclusion_explanation_uses_readable_category_names(self):
        _, _, relations = cap_review(
            [
                "Each party's aggregate liability will not exceed the Fees paid in the preceding twelve months.",
                "Indemnification obligations are excluded from the foregoing cap.",
                "Indemnification liability is capped at $1,000,000.",
            ]
        )

        explanation = relations[0]["explanation"]
        self.assertIn("separate limit", explanation)
        self.assertIn("indemnification", explanation.lower())
        self.assertNotIn(r"indemn\w*", explanation)

    def test_multiple_same_block_exceptions_keep_the_material_indemnity_scope(self):
        document, check, relations = cap_review(
            [
                "Aggregate liability will not exceed fees paid in the preceding twelve months.",
                "The liability cap does not apply to fraud or willful misconduct. "
                "The liability cap does not apply to indemnification obligations.",
                "Indemnification liability is capped at $1,000,000.",
            ]
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertEqual(1, len(relations))
        self.assertEqual("indemnification obligations", relations[0]["exception_scope"])
        self.assertTrue(relations[0]["separate_limit_located"])
        cap_issue = next(issue for issue in analyze(document, CAP_ONLY)[0]
                         if issue["rule_id"] == "liability-cap")
        self.assertIn("Reconcile the indemnification obligations carveout", cap_issue["negotiating_response"])
        self.assertIn("amount, period, covered claims", cap_issue["negotiating_response"])

    def test_named_exclusion_is_not_suppressed_by_narrow_carveout_in_same_block(self):
        _, check, relations = cap_review(
            [
                "Aggregate liability will not exceed fees paid in the preceding twelve months.",
                "The liability cap does not apply to fraud or willful misconduct. "
                "Indemnification obligations are excluded from the foregoing cap.",
            ]
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertEqual(1, len(relations))
        self.assertEqual("indemnification", relations[0]["exception_scope"])

    def test_multiple_named_exclusions_in_same_block_are_all_retained(self):
        _, check, relations = cap_review(
            [
                "Aggregate liability will not exceed fees paid in the preceding twelve months.",
                "Indemnification obligations are excluded from the foregoing cap. "
                "Confidentiality obligations are excluded from the foregoing cap.",
            ]
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertEqual({"indemnification", "confidentiality"},
                         {relation["exception_scope"] for relation in relations})

    def test_inline_except_for_named_category_is_linked_to_the_cap(self):
        document = extract(
            b"Aggregate liability is capped at Fees paid in the prior twelve months, "
            b"except for indemnification obligations."
        )
        issues, checks = analyze(document, CAP_ONLY)
        cap_issue = next(issue for issue in issues if issue["rule_id"] == "liability-cap")

        self.assertEqual("NEEDS_REVIEW", checks[0]["outcome"])
        self.assertEqual("cap-exception", cap_issue["interactions"][0]["kind"])
        self.assertIn("State whether indemnification", cap_issue["negotiating_response"])
        self.assertTrue(all(document["text"][a["start"]:a["end"]] == a["text"]
                            for a in cap_issue["interactions"][0]["evidence"]))

    def test_inline_negated_exception_is_not_reported_as_a_cap_relation(self):
        _, check, relations = cap_review(
            ["Aggregate liability will not exceed Fees paid in the preceding twelve months, "
             "and the cap does not exclude indemnification obligations."]
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertEqual([], relations)

    def test_negated_exclusion_does_not_create_a_cap_exception_relation(self):
        _, check, relations = cap_review(
            ["Aggregate liability will not exceed Fees paid in the preceding twelve months.",
             "Indemnification obligations are not excluded from the foregoing cap."]
        )

        self.assertEqual("COMPLIANT", check["outcome"])
        self.assertEqual([], relations)

    def test_inline_exception_does_not_borrow_cap_from_a_prior_sentence(self):
        _, check, relations = cap_review(
            ["Aggregate liability will not exceed Fees paid in the preceding twelve months. "
             "Except for indemnification obligations, a notice of claim must be provided promptly."]
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertEqual([], relations)


class TermExpirationRecognition(unittest.TestCase):
    def review(self, text):
        document = extract(text.encode())
        issues, checks = analyze(document, NDA_PLAYBOOK)
        check = next(check for check in checks if check["rule_id"] == "term-expiration")
        findings = [issue for issue in issues if issue["rule_id"] == "term-expiration"]
        return document, check, findings

    def test_held_out_agreement_expiry_is_found_and_spelled_duration_is_recognized(self):
        document, check, findings = self.review(
            "This agreement expires two years after signing. "
            "Confidentiality obligations survive for three years after termination."
        )

        self.assertEqual("COMPLIANT", check["outcome"])
        self.assertEqual([], findings)
        self.assertEqual("two years", document["blocks"][0]["text"].split("expires ", 1)[1].split(" after", 1)[0])

    def test_established_begins_and_continues_term_remains_supported(self):
        _, check, findings = self.review(
            "This agreement begins on the effective date and continues for twelve months, "
            "then expires without automatic renewal."
        )

        self.assertEqual("COMPLIANT", check["outcome"])
        self.assertEqual([], findings)

    def test_agreement_expiry_in_one_paragraph_is_separate_from_survival_period(self):
        _, check, findings = self.review(
            "This agreement will end two years from the date both parties sign.\n\n"
            "Confidentiality obligations survive for three years after termination."
        )

        self.assertEqual("COMPLIANT", check["outcome"])
        self.assertEqual([], findings)

    def test_established_initial_term_wording_remains_supported(self):
        _, check, findings = self.review("The initial term is 24 months.")

        self.assertEqual("COMPLIANT", check["outcome"])
        self.assertEqual([], findings)

    def test_compound_written_duration_is_compared_as_one_number(self):
        _, check, findings = self.review("This agreement expires in twenty-one months.")

        self.assertEqual("COMPLIANT", check["outcome"])
        self.assertEqual([], findings)

    def test_custom_duration_rule_cannot_borrow_survival_period_from_other_block(self):
        term = next(rule for rule in NDA_PLAYBOOK.rules if rule.id == "term-expiration")
        custom_term = term.model_copy(update={"acceptable_all": [[r"\b3 years\b"]]})
        custom_book = NDA_PLAYBOOK.model_copy(update={
            "rules": [custom_term if rule.id == "term-expiration" else rule for rule in NDA_PLAYBOOK.rules]
        })
        document = extract(
            b"This agreement expires for 2 years after signing.\n\n"
            b"Confidentiality obligations survive for 3 years after termination."
        )

        issues, checks = analyze(document, custom_book)
        check = next(check for check in checks if check["rule_id"] == "term-expiration")
        findings = [issue for issue in issues if issue["rule_id"] == "term-expiration"]

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertEqual(document["blocks"][0], findings[0]["source_anchor"])

    def test_custom_duration_rule_cannot_borrow_survival_period_from_same_block(self):
        term = next(rule for rule in NDA_PLAYBOOK.rules if rule.id == "term-expiration")
        custom_term = term.model_copy(update={"acceptable_all": [[r"\b3 years\b"]]})
        custom_book = NDA_PLAYBOOK.model_copy(update={
            "rules": [custom_term if rule.id == "term-expiration" else rule for rule in NDA_PLAYBOOK.rules]
        })
        document = extract(
            b"This agreement expires after 2 years. Confidentiality obligations survive for 3 years."
        )

        issues, checks = analyze(document, custom_book)
        check = next(check for check in checks if check["rule_id"] == "term-expiration")
        findings = [issue for issue in issues if issue["rule_id"] == "term-expiration"]

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertEqual(document["blocks"][0], findings[0]["source_anchor"])

    def test_conflicting_agreement_term_sentences_cannot_satisfy_strict_duration(self):
        term = next(rule for rule in NDA_PLAYBOOK.rules if rule.id == "term-expiration")
        custom_term = term.model_copy(update={"acceptable_all": [[r"\b3 years\b"]]})
        custom_book = NDA_PLAYBOOK.model_copy(update={
            "rules": [custom_term if rule.id == "term-expiration" else rule for rule in NDA_PLAYBOOK.rules]
        })
        document = extract(
            b"This agreement expires after 2 years. This agreement will end after 3 years."
        )

        issues, checks = analyze(document, custom_book)
        check = next(check for check in checks if check["rule_id"] == "term-expiration")
        findings = [issue for issue in issues if issue["rule_id"] == "term-expiration"]

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertEqual(document["blocks"][0], findings[0]["source_anchor"])

    def test_conflicting_default_term_durations_across_sentences_require_review(self):
        _, check, findings = self.review(
            "This agreement expires after 12 months. This agreement will end after 24 months."
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertFalse(findings[0]["absence"])

    def test_conflicting_default_term_durations_across_blocks_require_review(self):
        _, check, findings = self.review(
            "This agreement expires after 12 months.\n\n"
            "This agreement will end after 24 months."
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertFalse(findings[0]["absence"])

    def test_conflicting_expiration_dates_across_sentences_require_review(self):
        _, check, findings = self.review(
            "This agreement expires on December 31, 2028. "
            "This agreement terminates on January 31, 2029."
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertFalse(findings[0]["absence"])

    def test_conflicting_expiration_dates_across_blocks_require_review(self):
        _, check, findings = self.review(
            "This agreement expires on December 31, 2028.\n\n"
            "This agreement terminates on January 31, 2029."
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertFalse(findings[0]["absence"])

    def test_repeated_identical_term_durations_are_consistent(self):
        _, check, findings = self.review(
            "This agreement expires after 12 months. This agreement ends after 12 months."
        )

        self.assertEqual("COMPLIANT", check["outcome"])
        self.assertEqual([], findings)

    def test_equivalent_month_and_year_durations_remain_unresolved(self):
        _, check, findings = self.review(
            "This agreement expires after 12 months. This agreement ends after one year."
        )

        # The terms may be equivalent, but the configured policy pattern does
        # not accept singular "year" wording; do not assert compliance.
        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertFalse(findings[0]["absence"])

    def test_survival_period_alone_does_not_satisfy_agreement_term(self):
        _, check, findings = self.review(
            "Confidentiality obligations survive for three years after termination."
        )

        self.assertEqual("ISSUE", check["outcome"])
        self.assertEqual(1, len(findings))
        self.assertTrue(findings[0]["absence"])

    def test_negated_expiry_stays_located_but_unresolved(self):
        document, check, findings = self.review(
            "This agreement does not expire two years after signing."
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertEqual(1, len(findings))
        self.assertFalse(findings[0]["absence"])
        self.assertEqual(document["blocks"][0], findings[0]["source_anchor"])
        self.assertIn("negation", findings[0]["explanation"])

    def test_conditional_expiry_stays_unresolved(self):
        _, check, findings = self.review(
            "This agreement expires two years after signing unless extended in writing."
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertFalse(findings[0]["absence"])

    def test_subordinate_confidentiality_expiry_does_not_define_agreement_term(self):
        _, check, findings = self.review(
            "This agreement requires confidentiality obligations to expire after three years."
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertFalse(findings[0]["absence"])

    def test_renewal_period_alone_does_not_establish_agreement_term(self):
        _, check, findings = self.review("Renewal term is two years.")

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertFalse(findings[0]["absence"])

    def test_expiration_date_remains_supported(self):
        _, check, findings = self.review("This agreement expires on December 31, 2028.")

        self.assertEqual("COMPLIANT", check["outcome"])
        self.assertEqual([], findings)

    def test_condition_in_the_next_paragraph_prevents_compliance(self):
        _, check, findings = self.review(
            "This agreement expires two years after signing.\nUnless extended in writing."
        )

        self.assertEqual("NEEDS_REVIEW", check["outcome"])
        self.assertFalse(findings[0]["absence"])

    def test_narrow_early_termination_qualifiers_preserve_definite_base_term(self):
        for text in (
            "The term of this agreement is two years, subject to early termination.",
            "This agreement expires after two years unless terminated earlier under this agreement.",
            "This agreement expires after two years unless ended earlier under the agreement.",
        ):
            with self.subTest(text=text):
                _, check, findings = self.review(text)
                self.assertEqual("COMPLIANT", check["outcome"])
                self.assertEqual([], findings)

    def test_early_termination_exception_does_not_accept_extension_or_unbounded_condition(self):
        for text in (
            "This agreement expires after two years unless extended in writing.",
            "This agreement expires after two years unless terminated for convenience.",
            "This agreement expires after two years unless terminated earlier by written notice.",
            "This agreement expires after two years subject to renewal by either party.",
            "This agreement expires after two years subject to early termination by either party.",
        ):
            with self.subTest(text=text):
                _, check, findings = self.review(text)
                self.assertEqual("NEEDS_REVIEW", check["outcome"])
                self.assertFalse(findings[0]["absence"])


class ExtractionLayoutRegression(unittest.TestCase):
    @settings(max_examples=120, deadline=None)
    @given(st.binary(max_size=2048))
    def test_arbitrary_small_inputs_fail_safely_or_return_exact_spans(self, payload):
        try:
            extracted = extract(payload)
        except ParseError:
            return

        self.assertEqual(
            "\n\n".join(block["text"] for block in extracted["blocks"]),
            extracted["text"],
        )
        self.assertTrue(
            all(
                extracted["text"][block["start"] : block["end"]] == block["text"]
                for block in extracted["blocks"]
            )
        )

    def test_nested_docx_table_text_is_retained_in_document_order(self):
        source = Document()
        source.add_paragraph("Before table.")
        outer = source.add_table(rows=1, cols=1)
        cell = outer.cell(0, 0)
        cell.paragraphs[0].text = "Outer table term."
        nested = cell.add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Nested table exception."
        source.add_paragraph("After table.")
        payload = BytesIO()
        source.save(payload)

        extracted = extract(payload.getvalue())

        self.assertEqual(
            ["Before table.", "Outer table term.", "Nested table exception.", "After table."],
            [block["text"] for block in extracted["blocks"]],
        )

    def test_hidden_character_style_cannot_enter_analysis_as_visible_text(self):
        source = Document()
        base = source.styles.add_style("HiddenBase", WD_STYLE_TYPE.CHARACTER)
        base.font.hidden = True
        hidden = source.styles.add_style("HiddenClause", WD_STYLE_TYPE.CHARACTER)
        hidden.base_style = base
        paragraph = source.add_paragraph()
        paragraph.add_run("Visible provision. ")
        secret = paragraph.add_run("Hidden liability exception.")
        secret.style = hidden
        payload = BytesIO()
        source.save(payload)

        with self.assertRaisesRegex(ParseError, "hidden text"):
            extract(payload.getvalue())

    def test_hidden_default_run_formatting_cannot_enter_analysis(self):
        source = Document()
        defaults = source.styles.element.find(qn("w:docDefaults"))
        run_defaults = defaults.find(qn("w:rPrDefault"))
        run_properties = run_defaults.find(qn("w:rPr"))
        if run_properties is None:
            run_properties = OxmlElement("w:rPr")
            run_defaults.append(run_properties)
        run_properties.append(OxmlElement("w:vanish"))
        source.add_paragraph("Text rendered hidden by document defaults.")
        payload = BytesIO()
        source.save(payload)

        with self.assertRaisesRegex(ParseError, "hidden text"):
            extract(payload.getvalue())

    def test_hidden_default_paragraph_and_character_styles_are_detected_without_style_refs(self):
        paragraph_default = Document()
        paragraph_default.styles["Normal"].font.hidden = True
        paragraph_default.add_paragraph("Hidden by default paragraph style.")
        paragraph_payload = BytesIO()
        paragraph_default.save(paragraph_payload)

        character_default = Document()
        character_default.styles["Default Paragraph Font"].font.hidden = True
        character_default.add_paragraph("Hidden by default character style.")
        character_payload = BytesIO()
        character_default.save(character_payload)

        for payload in (paragraph_payload.getvalue(), character_payload.getvalue()):
            with self.subTest(payload=payload[:24]), self.assertRaisesRegex(ParseError, "hidden text"):
                extract(payload)

    def test_pdf_layout_pair_budget_is_never_exceeded(self):
        from workbench.parsing import _inspect_pdf_layout

        blocks = [
            (0.0, 0.0, 10.0, 10.0, "text", 0, 0)
            for _ in range(450)
        ]

        found, comparisons, limited = _inspect_pdf_layout(blocks, 100_000)

        self.assertFalse(found)
        self.assertLessEqual(comparisons, 100_000)
        self.assertTrue(limited)

    def test_docx_comments_are_disclosed_as_not_analyzed(self):
        source = Document()
        source.add_paragraph("Visible operative text.")
        payload = BytesIO()
        source.save(payload)

        commented = BytesIO()
        with ZipFile(BytesIO(payload.getvalue())) as original, ZipFile(commented, "w") as updated:
            for entry in original.infolist():
                updated.writestr(entry, original.read(entry.filename))
            updated.writestr(
                "word/comments.xml",
                b'<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                b'<w:comment w:id="0"><w:p><w:r><w:t>Reviewer comment contains an exception.</w:t></w:r></w:p></w:comment>'
                b'</w:comments>',
            )

        extracted = extract(commented.getvalue())

        self.assertTrue(any("comment" in warning.lower() for warning in extracted["warnings"]))

    @settings(max_examples=60, deadline=None)
    @given(
        st.lists(
            st.text(alphabet="abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ ,.-", min_size=1, max_size=100).filter(str.strip),
            min_size=1,
            max_size=12,
        )
    )
    def test_docx_extraction_roundtrips_generated_paragraphs(self, paragraphs):
        source = Document()
        for paragraph in paragraphs:
            source.add_paragraph(paragraph)
        payload = BytesIO()
        source.save(payload)

        extracted = extract(payload.getvalue())

        self.assertEqual(paragraphs, [block["text"] for block in extracted["blocks"]])
        self.assertEqual("\n\n".join(paragraphs), extracted["text"])
        self.assertTrue(
            all(
                extracted["text"][block["start"] : block["end"]] == block["text"]
                for block in extracted["blocks"]
            )
        )

    def test_docx_table_terms_keep_body_order_and_exact_offsets(self):
        source = Document()
        source.add_paragraph("Opening provision.")
        table = source.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Left cell term."
        table.cell(0, 1).text = "Right cell term."
        source.add_paragraph("Closing provision.")
        payload = BytesIO()
        source.save(payload)

        extracted = extract(payload.getvalue())

        self.assertEqual(
            ["Opening provision.", "Left cell term.", "Right cell term.", "Closing provision."],
            [block["text"] for block in extracted["blocks"]],
        )
        self.assertTrue(
            all(
                extracted["text"][block["start"] : block["end"]] == block["text"]
                for block in extracted["blocks"]
            )
        )

    def test_pdf_page_blocks_retain_page_identity_and_exact_offsets(self):
        source = pymupdf.open()
        for text in ("First page condition.", "Second page exception."):
            page = source.new_page()
            page.insert_text((72, 72), text)
        payload = source.tobytes()
        source.close()

        extracted = extract(payload)

        self.assertEqual([1, 2], [block["page"] for block in extracted["blocks"]])
        self.assertEqual(
            ["First page condition.", "Second page exception."],
            [block["text"] for block in extracted["blocks"]],
        )
        self.assertTrue(
            all(
                extracted["text"][block["start"] : block["end"]] == block["text"]
                for block in extracted["blocks"]
            )
        )

    def test_pdf_side_by_side_text_warns_that_order_needs_visual_review(self):
        source = pymupdf.open()
        page = source.new_page(width=612, height=792)
        page.insert_textbox(
            pymupdf.Rect(40, 60, 280, 500),
            "Section 4. Liability\nEach party liability will not exceed 100.\nIndemnification applies to third-party claims.",
            fontsize=12,
        )
        page.insert_textbox(
            pymupdf.Rect(330, 60, 570, 500),
            "Section 5. Other Terms\nPayment is due in 30 days.\nConfidentiality lasts five years.",
            fontsize=12,
        )
        payload = source.tobytes()
        source.close()

        extracted = extract(payload)

        self.assertTrue(any("side-by-side" in warning.lower() for warning in extracted["warnings"]))

    def test_pdf_annotations_are_disclosed_as_not_analyzed(self):
        source = pymupdf.open()
        page = source.new_page()
        page.insert_text((72, 72), "Visible operative text.")
        page.add_text_annot((72, 100), "Reviewer comment contains a proposed exception.")
        payload = source.tobytes()
        source.close()

        extracted = extract(payload)

        self.assertTrue(any("annotation" in warning.lower() for warning in extracted["warnings"]))


class ReviewRevisionRegression(unittest.TestCase):
    def test_old_review_is_labeled_without_rewriting_its_decision(self):
        stored = {
            "review_revision": REVIEW_REVISION - 1,
            "extraction_revision": EXTRACTION_REVISION,
            "issues": [{"rule_id": "liability-cap", "status": "accepted", "notes": "Existing note."}],
            "playbook_snapshot": {"rules": []},
        }
        original = deepcopy(stored)

        visible = review_view(stored)

        self.assertEqual("HISTORICAL_UNVERIFIED", visible["verification_status"])
        self.assertIn("current recognition and clause-interaction checks", visible["verification_warning"])
        self.assertEqual("accepted", visible["issues"][0]["status"])
        self.assertEqual("Existing note.", visible["issues"][0]["notes"])
        self.assertEqual(original, stored)

    def test_older_extraction_requires_reupload_warning(self):
        visible = review_view(
            {
                "review_revision": REVIEW_REVISION,
                "extraction_revision": EXTRACTION_REVISION - 1,
                "issues": [],
                "playbook_snapshot": {"rules": []},
            }
        )

        self.assertEqual("HISTORICAL_UNVERIFIED", visible["verification_status"])
        self.assertIn("retained original", visible["verification_warning"])

    def test_current_revision_does_not_show_historical_warning(self):
        visible = review_view(
            {
                "review_revision": REVIEW_REVISION,
                "extraction_revision": EXTRACTION_REVISION,
                "issues": [],
                "playbook_snapshot": {"rules": []},
            }
        )

        self.assertEqual("CURRENT_CHECKS", visible["verification_status"])
        self.assertIsNone(visible["verification_warning"])

if __name__ == "__main__":
    unittest.main()
