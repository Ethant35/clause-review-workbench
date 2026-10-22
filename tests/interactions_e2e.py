"""Fixed, real HTTP/parser/SQLite/export/browser checks for cross-clause analysis.

Author: Ethan. No engine imports, mocks, or model calls.
"""
import argparse
from datetime import datetime, timezone
from io import BytesIO
import json
import os
from pathlib import Path
import signal
import socket
import subprocess
import sys
import tempfile
import time

import httpx
from docx import Document
import pymupdf
from playwright.sync_api import sync_playwright, expect
import yaml

ROOT = Path(__file__).resolve().parents[1]


def encoded(blocks, fmt):
    if fmt == 'txt':
        return '\n\n'.join(blocks).encode()
    if fmt == 'docx':
        doc = Document()
        doc.core_properties.author = doc.core_properties.last_modified_by = 'Ethan'
        for block in blocks:
            doc.add_paragraph(block)
        out = BytesIO(); doc.save(out)
        return out.getvalue()
    doc = pymupdf.open()
    for block in blocks:
        # One paragraph per page keeps independent blocks and exercises page links.
        page = doc.new_page()
        text = block.replace('’', "'").replace('“', '"').replace('”', '"')
        assert page.insert_textbox(pymupdf.Rect(50, 50, 550, 740), text, fontsize=11) >= 0
    doc.set_metadata({'author': 'Ethan'})
    out = doc.tobytes(); doc.close()
    return out


def canonical(text):
    return ' '.join(text.replace('’', "'").replace('“', '"').replace('”', '"').split())


def run(output, browser_check):
    out = ROOT / 'artifacts'; out.mkdir(exist_ok=True)
    run_dir = Path(tempfile.mkdtemp(prefix='interactions-', dir=out))
    network = run_dir / 'network.jsonl'; network.touch()
    with socket.socket() as sock:
        sock.bind(('127.0.0.1', 0)); port = sock.getsockname()[1]
    base = f'http://127.0.0.1:{port}'
    env = os.environ.copy(); env.pop('OPENROUTER_API_KEY', None)
    env.update(CRW_DATA_DIR=str(run_dir/'data'), CRW_PORT=str(port), CRW_NETWORK_AUDIT=str(network))
    log = (run_dir/'server.log').open('w')
    server = subprocess.Popen([sys.executable, 'tests/serve_checked.py'], cwd=ROOT, env=env,
                              stdout=log, stderr=subprocess.STDOUT, start_new_session=True)
    registry = Path(os.environ.get('CRW_PROCESS_REGISTRY', out/'processes.json'))
    registry.write_text(json.dumps([os.getpgrp(), server.pid]))
    results = []; example = None
    try:
        with httpx.Client(base_url=base, timeout=40, trust_env=False) as client:
            for _ in range(100):
                try:
                    if client.get('/api/health').status_code == 200: break
                except httpx.HTTPError: pass
                time.sleep(.1)
            manifest = json.loads((ROOT/'corpus/interactions/cases.json').read_text())
            pb = yaml.safe_load(client.get('/api/playbooks/commercial').json()['yaml'])
            policies = {}
            for rule_id in {c['rule_id'] for c in manifest['cases']}:
                focused = {**pb, 'name': 'Interaction acceptance '+rule_id, 'rules': [r for r in pb['rules'] if r['id'] == rule_id]}
                response = client.post('/api/playbooks', json={'yaml': yaml.safe_dump(focused)})
                response.raise_for_status(); policies[rule_id] = response.json()['id']
            for case in manifest['cases']:
                for fmt in ('txt', 'docx', 'pdf'):
                    response = client.post('/api/documents', files={'file': (case['id']+'.'+fmt, encoded(case['blocks'], fmt))})
                    response.raise_for_status(); doc = response.json()
                    response = client.post(f"/api/documents/{doc['id']}/reviews", json={'playbook_id': policies[case['rule_id']]})
                    response.raise_for_status(); review = response.json()
                    relations = [x for i in review['issues'] for x in i.get('interactions', [])]
                    evidence = [b for x in relations for b in x['evidence']]
                    actual_evidence = {canonical(b['text']) for b in evidence}
                    expected_evidence = {canonical(case['blocks'][i]) for i in case.get('expected_evidence', [])}
                    checks = {
                        'extraction': [canonical(b['text']) for b in doc['blocks']] == [canonical(b) for b in case['blocks']],
                        'kinds': set(x['kind'] for x in relations) == set(case['expected_kinds']),
                        'outcome': review['checks'][0]['outcome'] == case['expected_outcome'],
                        'evidence_complete': actual_evidence == expected_evidence,
                        'evidence_exact': all(b in doc['blocks'] and doc['text'][b['start']:b['end']] == b['text'] for b in evidence),
                        'status': not case.get('expected_statuses') or sorted(x.get('status') for x in relations) == sorted(case['expected_statuses']),
                        'explanation': not case.get('explanation_contains') or any(case['explanation_contains'] in x['explanation'] for x in relations),
                        'uncertainty': all(i['confidence']=='NEEDS_REVIEW' for i in review['issues'] if i.get('interactions')),
                    }
                    # Persistence and all export formats must carry the same relationships.
                    saved = client.get('/api/reviews/'+review['id']).json()
                    checks['storage'] = saved['issues'] == review['issues']
                    for export_fmt, extension in [('json','json'),('markdown','md'),('docx','docx')]:
                        exported = client.get(f"/api/reviews/{review['id']}/export", params={'format': export_fmt})
                        exported.raise_for_status()
                        if export_fmt == 'docx':
                            document = Document(BytesIO(exported.content))
                            text = '\n'.join(p.text for p in document.paragraphs)
                            checks['author'] = document.core_properties.author == document.core_properties.last_modified_by == 'Ethan'
                        elif export_fmt == 'json':
                            text = json.dumps(exported.json(), ensure_ascii=False)
                            checks['json_identity'] = exported.json()['issues'] == review['issues']
                        else: text = exported.text
                        # Markdown escapes text; the explanation and every evidence location remain readable.
                        checks['export_'+export_fmt] = all(x['explanation'] in text and all(b['location'] in text for b in x['evidence']) for x in relations)
                        if case['id']=='cap-exception-supercap' and fmt=='docx':
                            (run_dir/('interaction-report.'+extension)).write_bytes(exported.content)
                    row = {'id':case['id'], 'format':fmt, 'actual_outcome':review['checks'][0]['outcome'], 'actual_kinds':[x['kind'] for x in relations], 'checks':checks, 'pass':all(checks.values())}
                    results.append(row)
                    if case['id']=='cap-exception-supercap' and fmt=='docx': example = (doc,review)
            for label, blocks in [
                ('interaction-count-limit', [manifest['cases'][0]['blocks'][0]] * 101 + ['The liability cap does not apply to indemnification obligations.']),
                ('reference-evidence-limit', ['Subject to Section 9, '+manifest['cases'][0]['blocks'][0], 'Section 9. Conditions'] + ['An additional condition applies.'] * 65),
            ]:
                response = client.post('/api/documents', files={'file':(label+'.txt', encoded(blocks,'txt'))})
                response.raise_for_status(); did=response.json()['id']
                response = client.post(f'/api/documents/{did}/reviews', json={'playbook_id':policies['liability-cap']})
                results.append({'id':label,'pass':response.status_code==422 and client.get(f'/api/documents/{did}/reviews').json()==[]})
            # Same agreement under NDA-only policy must not activate commercial rules.
            response=client.post(f"/api/documents/{example[0]['id']}/reviews",json={'playbook_id':'nda'})
            response.raise_for_status()
            results.append({'id':'inactive-rule-isolation','pass':not any(i.get('interactions') for i in response.json()['issues'])})
            # Restore the focused saved review as newest for the browser flow.
            response=client.post(f"/api/documents/{example[0]['id']}/reviews",json={'playbook_id':policies['liability-cap']})
            response.raise_for_status(); example=(example[0],response.json())
            if browser_check:
                doc, review = example
                issue = next(i for i in review['issues'] if i.get('interactions'))
                with sync_playwright() as pw:
                    browser = pw.chromium.launch()
                    page = browser.new_page(viewport={'width':1440,'height':1100})
                    page.goto(base)
                    page.locator('.doc-row').filter(has_text=doc['name']).click()
                    expect(page.get_by_role('heading', name='Related clauses and effect')).to_be_visible()
                    expect(page.locator('.interaction-card')).to_have_count(1)
                    expect(page.locator('.interaction-card .source-card')).to_have_count(3)
                    page.screenshot(path=str(run_dir/'interaction-finding.png'), full_page=True)
                    page.locator('.interaction-card').get_by_role('button',name='View in source').last.click()
                    target=issue['interactions'][0]['evidence'][-1]
                    expect(page.locator('#block-'+target['id'])).to_have_class('source-block source-highlight')
                    page.screenshot(path=str(run_dir/'related-source.png'), full_page=True)
                    page.get_by_role('tab',name='Findings',exact=False).click()
                    page.get_by_label('Reviewer notes').fill('Review separate indemnity limit with counsel.')
                    page.get_by_role('button',name='Save note',exact=True).click()
                    expect(page.get_by_text('Reviewer update saved locally.',exact=True)).to_be_visible()
                    page.reload()
                    page.locator('.doc-row').filter(has_text=doc['name']).click()
                    expect(page.get_by_label('Reviewer notes')).to_have_value('Review separate indemnity limit with counsel.')
                    expect(page.locator('.interaction-card')).to_have_count(1)
                    page.set_viewport_size({'width':390,'height':844})
                    expect(page.get_by_role('heading',name='Related clauses and effect')).to_be_visible()
                    assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
                    page.screenshot(path=str(run_dir/'interaction-mobile.png'),full_page=True)
                    browser.close()
                results.append({'id':'browser-evidence-navigation-and-notes','pass':True})
            results.append({'id':'offline-network','pass':network.read_text()==''})
    except Exception as exc:
        results.append({'id': 'run-error', 'pass': False, 'error': str(exc)[:2000]})
        raise
    finally:
        if server.poll() is None: os.killpg(server.pid,signal.SIGTERM); server.wait(timeout=10)
        registry.write_text('[]'); log.close()
        report={'author':'Ethan','timestamp':datetime.now(timezone.utc).isoformat(),'commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),'run':str(run_dir.relative_to(ROOT)),'passed':sum(x['pass'] for x in results),'total':len(results),'results':results}
        Path(output).write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps({k:report[k] for k in ('passed','total','run')}))
    for row in results:
        if not row['pass']: print('FAIL',row['id'],row.get('format'),[k for k,v in row['checks'].items() if not v],row['actual_outcome'])
    return all(x['pass'] for x in results)


if __name__ == '__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--output',default='artifacts/interactions-results.json');parser.add_argument('--no-browser',action='store_true')
    args=parser.parse_args(); raise SystemExit(0 if run(args.output,not args.no_browser) else 1)
