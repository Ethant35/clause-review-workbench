"""Real API counterexample suite. Expected meanings are authored independently in JSON."""
import argparse
from datetime import datetime, timezone
from io import BytesIO
import json
import os
from pathlib import Path
import signal
import socket
import subprocess
import sys
import tempfile
import time
from zipfile import ZipFile

import httpx
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
import pymupdf
import yaml

ROOT=Path(__file__).resolve().parents[1]

# Structural expectations are fixed before running the application. All test files are synthetic.
STRUCTURAL={
 'altchunk':('docx',422,'embedded'),
 'inline-control':('docx',422,'content control'),
 'moved-text':('docx',422,'revision'),
 'hidden-text':('docx',422,'hidden'),
 'hidden-style':('docx',422,'hidden'),
 'hidden-defaults':('docx',422,'hidden'),
 'hyperlink':('docx',201,None),
 'table':('docx',201,None),
 'raster-with-page-number':('pdf',422,'OCR'),
 'plain-pdf':('pdf',201,None),
}
TARGET='Provider may amend this agreement unilaterally without consent.'

def docx_bytes(text):
 d=Document();d.core_properties.author=d.core_properties.last_modified_by='Ethan'
 d.add_paragraph(text);f=BytesIO();d.save(f);return f.getvalue()

def pdf_bytes(text):
 d=pymupdf.open();p=d.new_page();assert p.insert_textbox(pymupdf.Rect(50,50,550,740),text,fontsize=11)>=0
 d.set_metadata({'author':'Ethan'});b=d.tobytes();d.close();return b

def special(kind):
 if kind=='plain-pdf':return pdf_bytes(TARGET)
 if kind=='raster-with-page-number':
  # A real rasterized clause with a live page number: nonempty extraction isn't completeness.
  source=pymupdf.open(stream=pdf_bytes(TARGET),filetype='pdf')
  raster=source[0].get_pixmap().tobytes('png');source.close()
  d=pymupdf.open();p=d.new_page();p.insert_image(p.rect,stream=raster);p.insert_text((30,800),'Page 1')
  d.set_metadata({'author':'Ethan'});b=d.tobytes();d.close();return b
 d=Document();d.core_properties.author=d.core_properties.last_modified_by='Ethan'
 if kind=='table':d.add_table(rows=1,cols=1).cell(0,0).text=TARGET
 elif kind=='hyperlink':
  p=d.add_paragraph();h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),'internal');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=TARGET;r.append(t);h.append(r);p._p.append(h)
 else:
  p=d.add_paragraph('The term of this agreement is twelve months.')
  if kind=='altchunk':
   element=OxmlElement('w:altChunk');element.set(qn('r:id'),'importedClause');d.element.body.insert(1,element)
  elif kind=='inline-control':
   s=OxmlElement('w:sdt');content=OxmlElement('w:sdtContent');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=TARGET;r.append(t);content.append(r);s.append(content);p._p.append(s)
  elif kind=='moved-text':
   m=OxmlElement('w:moveTo');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=TARGET;r.append(t);m.append(r);p._p.append(m)
  elif kind=='hidden-text':p.add_run(TARGET).font.hidden=True
  elif kind=='hidden-style':
   base=d.styles.add_style('HiddenBase',WD_STYLE_TYPE.CHARACTER);base.font.hidden=True
   style=d.styles.add_style('HiddenClause',WD_STYLE_TYPE.CHARACTER);style.base_style=base
   run=p.add_run(TARGET);run.style=style
  elif kind=='hidden-defaults':
   defaults=d.styles.element.find(qn('w:docDefaults'));run_defaults=defaults.find(qn('w:rPrDefault'));properties=run_defaults.find(qn('w:rPr'))
   if properties is None:properties=OxmlElement('w:rPr');run_defaults.append(properties)
   properties.append(OxmlElement('w:vanish'));p.add_run(TARGET)
 out=BytesIO();d.save(out)
 if kind!='altchunk':return out.getvalue()
 # Complete the real OOXML relationship/content type and imported text payload.
 output=BytesIO()
 with ZipFile(BytesIO(out.getvalue())) as src, ZipFile(output,'w') as dst:
  for info in src.infolist():
   data=src.read(info.filename)
   if info.filename=='word/_rels/document.xml.rels':
    root=etree.fromstring(data);rel=etree.SubElement(root,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship');rel.attrib.update({'Id':'importedClause','Type':'http://schemas.openxmlformats.org/officeDocument/2006/relationships/aFChunk','Target':'import.txt'});data=etree.tostring(root)
   if info.filename=='[Content_Types].xml':
    root=etree.fromstring(data);e=etree.SubElement(root,'{http://schemas.openxmlformats.org/package/2006/content-types}Override');e.attrib.update({'PartName':'/word/import.txt','ContentType':'text/plain'});data=etree.tostring(root)
   dst.writestr(info,data)
  dst.writestr('word/import.txt',TARGET)
 return output.getvalue()

def run(output):
 manifest=json.loads((ROOT/'corpus/adversarial/cases.json').read_text());results=[]
 out=ROOT/'artifacts';out.mkdir(exist_ok=True);run=Path(tempfile.mkdtemp(prefix='adversarial-',dir=out));network=run/'network.jsonl';network.touch()
 with socket.socket() as sock:sock.bind(('127.0.0.1',0));port=sock.getsockname()[1]
 env=os.environ.copy();env.pop('OPENROUTER_API_KEY',None);env.update(CRW_DATA_DIR=str(run/'data'),CRW_PORT=str(port),CRW_NETWORK_AUDIT=str(network))
 log=(run/'server.log').open('w');p=subprocess.Popen([sys.executable,'tests/serve_checked.py'],cwd=ROOT,env=env,stdout=log,stderr=subprocess.STDOUT,start_new_session=True)
 registry=Path(os.environ.get('CRW_PROCESS_REGISTRY',out/'processes.json'));registry.write_text(json.dumps([os.getpgrp(),p.pid]))
 try:
  with httpx.Client(base_url=f'http://127.0.0.1:{port}',timeout=40,trust_env=False) as c:
   for _ in range(100):
    try:
     if c.get('/api/health').status_code==200:break
    except httpx.HTTPError:pass
    time.sleep(.1)
   pb=yaml.safe_load(c.get('/api/playbooks/commercial').json()['yaml']);ids={}
   for ruleid in {x['rule_id'] for x in manifest['cases']}:
    policy={**pb,'name':'Independent focused acceptance '+ruleid,'rules':[r for r in pb['rules'] if r['id']==ruleid]}
    r=c.post('/api/playbooks',json={'yaml':yaml.safe_dump(policy)});r.raise_for_status();ids[ruleid]=r.json()['id']
   for case in manifest['cases']:
    for fmt in ['txt','docx','pdf']:
     data=case['text'].encode() if fmt=='txt' else docx_bytes(case['text']) if fmt=='docx' else pdf_bytes(case['text'])
     r=c.post('/api/documents',files={'file':(case['id']+'.'+fmt,data)});r.raise_for_status();doc=r.json()
     r=c.post(f"/api/documents/{doc['id']}/reviews",json={'playbook_id':ids[case['rule_id']]});r.raise_for_status();rev=r.json();check=rev['checks'][0]
     # Independent authored text is the extraction oracle; PDF wraps lines visually.
     exact=' '.join(doc['text'].split())==' '.join(case['text'].split())
     citations=all(not i['absence'] and i['source_anchor'] in doc['blocks'] and doc['text'][i['source_anchor']['start']:i['source_anchor']['end']]==i['source_anchor']['text'] and i['observed_language'] in i['source_anchor']['text'] for i in rev['issues'])
     confidence=all(i['confidence']=='NEEDS_REVIEW' for i in rev['issues']) if case['expected']=='NEEDS_REVIEW' else True
     properties=check['properties'].get('explicit_fee_cap')==case['explicit_fee_cap'] if 'explicit_fee_cap' in case else True
     row={'id':case['id'],'format':fmt,'expected':case['expected'],'actual':check['outcome'],'exact_extraction':exact,'valid_nonabsence_citations':citations,'correct_uncertainty':confidence,'correct_properties':properties}
     row['pass']=check['outcome']==case['expected'] and exact and citations and confidence and properties;results.append(row)
     c.delete('/api/documents/'+doc['id']).raise_for_status()
   for kind,(fmt,status,message) in STRUCTURAL.items():
    data=special(kind);(run/(kind+'.'+fmt)).write_bytes(data)
    r=c.post('/api/documents',files={'file':(kind+'.'+fmt,data)})
    correct=r.status_code==status
    if correct and status==422:correct=message.casefold() in r.json()['detail'].casefold()
    if status==201 and r.status_code==201:
     doc=r.json();correct=correct and TARGET==' '.join(doc['text'].split()) and c.get('/api/documents/'+doc['id']+'/original').content==data
    results.append({'id':kind,'format':fmt,'expected_http':status,'actual_http':r.status_code,'pass':correct})
   results.append({'id':'offline-network','pass':network.read_text()==''})
 finally:
  if p.poll() is None:os.killpg(p.pid,signal.SIGTERM);p.wait(timeout=10)
  registry.write_text('[]');log.close()
 report={'author':'Ethan','timestamp':datetime.now(timezone.utc).isoformat(),'commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),'run':str(run.relative_to(ROOT)),'passed':sum(x['pass'] for x in results),'total':len(results),'results':results}
 Path(output).write_text(json.dumps(report,indent=2)+'\n')
 print(json.dumps({k:report[k] for k in ['passed','total','run']},indent=2))
 for r in results:
  if not r['pass']:print('FAIL',r)
 return report['passed']==report['total']

if __name__=='__main__':
 parser=argparse.ArgumentParser();parser.add_argument('--output',default='artifacts/adversarial-results.json');args=parser.parse_args();raise SystemExit(0 if run(args.output) else 1)
