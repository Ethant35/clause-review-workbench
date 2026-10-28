"""Real browser/HTTP/parser/database/export acceptance. No mocks or fake models."""
from contextlib import contextmanager
from datetime import datetime, timezone
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import signal
import socket
import sqlite3
import subprocess
import sys
import time
import traceback
from urllib.parse import urlparse
import uuid
from zipfile import ZipFile

import httpx
from docx import Document
import pymupdf
from playwright.sync_api import sync_playwright, expect
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from scripts.corpus_eval import evaluate
from scripts.make_fixtures import make

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts'
OUT.mkdir(exist_ok=True)
RUN=OUT/('run-'+uuid.uuid4().hex[:8])
RUN.mkdir()
DATA=RUN/'data'
NETWORK=RUN/'network.jsonl'
NETWORK.touch()
results=[]
process=None
pids=[]
REGISTRY=Path(os.environ.get('CRW_PROCESS_REGISTRY',str(OUT/'processes.json')))
with socket.socket() as sock:
    sock.bind(('127.0.0.1',0)); PORT=sock.getsockname()[1]
BASE=f'http://127.0.0.1:{PORT}'
client=httpx.Client(base_url=BASE, timeout=45, trust_env=False)

def register(): REGISTRY.write_text(json.dumps([os.getpgrp(), *pids]))
def stop():
    global process
    if process and process.poll() is None:
        os.killpg(process.pid,signal.SIGTERM)
        try: process.wait(timeout=10)
        except subprocess.TimeoutExpired:
            os.killpg(process.pid,signal.SIGKILL); process.wait()
    if process and process.pid in pids: pids.remove(process.pid)
    register()

def start(key=None, data=DATA):
    global process
    stop()
    env=os.environ.copy()
    env.pop('OPENROUTER_API_KEY',None)
    if key: env['OPENROUTER_API_KEY']=key
    env.update(CRW_DATA_DIR=str(data),CRW_PORT=str(PORT),CRW_NETWORK_AUDIT=str(NETWORK))
    process=subprocess.Popen([sys.executable,'tests/serve_checked.py'],cwd=ROOT,env=env,stdout=(RUN/'server.log').open('a'),stderr=subprocess.STDOUT,start_new_session=True)
    pids.append(process.pid); register()
    for _ in range(100):
        try:
            if client.get('/api/health').status_code==200: return
        except httpx.HTTPError: pass
        time.sleep(.1)
    raise RuntimeError('Server did not start: '+(RUN/'server.log').read_text()[-2000:])

def check(code, fn):
    try:
        detail=fn() or 'Verified against running application'
        results.append(dict(id=code,status='PASS',detail=detail))
        print(code,'PASS',flush=True)
    except Exception as exc:
        results.append(dict(id=code,status='FAIL',detail=str(exc),trace=traceback.format_exc()))
        print(code,'FAIL',str(exc),flush=True)
    (OUT/'e2e-results.json').write_text(json.dumps({'timestamp':datetime.now(timezone.utc).isoformat(),'run':str(RUN.relative_to(ROOT)),'results':results},indent=2)+'\n')

def upload(path=None,data=None,name='unseen.txt'):
    if path: name=path.name; data=path.read_bytes()
    r=client.post('/api/documents',files={'file':(name,data,'application/octet-stream')}); r.raise_for_status(); return r.json()
def review(doc,pb='commercial'):
    r=client.post(f"/api/documents/{doc['id']}/reviews",json={'playbook_id':pb}); r.raise_for_status(); return r.json()
def getrev(id):
    r=client.get('/api/reviews/'+id);r.raise_for_status();return r.json()
def byrule(rev,id): return next(i for i in rev['issues'] if i['rule_id']==id)
def validate(doc,rev):
    for issue in rev['issues']:
        a=issue['source_anchor']
        if issue['absence']: assert a is None and issue['observed_language'] is None
        else: assert a in doc['blocks'] and doc['text'][a['start']:a['end']]==a['text'] and issue['observed_language'] in a['text']

fixtures=make()
state={}
try:
    start()
    with sync_playwright() as pw:
        browser=pw.chromium.launch()
        context=browser.new_context(viewport={'width':1440,'height':1100})
        page=context.new_page()
        external=[]
        page.on('request',lambda r: external.append(r.url) if urlparse(r.url).hostname not in {'127.0.0.1','localhost',None} else None)
        page.goto(BASE)
        expect(page.get_by_role('heading',name='Document library')).to_be_visible()
        def ingest_docx():
            doc=upload(fixtures/'synthetic.docx');state['docx']=doc
            expected=[p.text for p in Document(fixtures/'synthetic.docx').paragraphs if p.text.strip()]
            assert [b['text'] for b in doc['blocks']]==expected
            assert client.get(f"/api/documents/{doc['id']}/original").content==(fixtures/'synthetic.docx').read_bytes()
            validate(doc,review(doc))
        check('CR-02',ingest_docx)
        def ingest_pdf():
            doc=upload(fixtures/'synthetic.pdf');state['pdf']=doc
            assert {b['page'] for b in doc['blocks']}=={1,2}
            with pymupdf.open(fixtures/'synthetic.pdf') as pdf:
                for b in doc['blocks']: assert b['text'] in pdf[b['page']-1].get_text()
            rev=review(doc);validate(doc,rev)
            assert {i['rule_id'] for i in rev['issues']}=={'automatic-renewal','audit-rights'}
        check('CR-03',ingest_pdf)
        def ingest_text():
            path=ROOT/'corpus/contracts/04-renewal-and-unilateral-audit.txt'
            page.locator('.upload-box input').set_input_files(path)
            expect(page.get_by_role('heading',level=2,name=path.name)).to_be_visible()
            page.get_by_label('Review against').select_option('commercial')
            page.get_by_role('button',name='Run review',exact=True).click()
            expect(page.locator('.issue-card')).to_have_count(2)
            docs=client.get('/api/documents').json();doc=next(d for d in docs if d['name']==path.name)
            doc=client.get('/api/documents/'+doc['id']).json()
            assert [b['text'] for b in doc['blocks']]==[x for x in path.read_text().splitlines() if x.strip()]
            state['text']=doc;state['review']=client.get(f"/api/documents/{doc['id']}/reviews").json()[0]
        check('CR-04',ingest_text)
        def invalid():
            bad=[(b'\x00\x01\xff','file.txt'),(b'%PDF-broken','bad.pdf'),(b'PKgarbage','bad.docx'),(b'','empty.txt'),(b'<html><script>alert(1)</script></html>','fake.txt'),((fixtures/'image-only.pdf').read_bytes(),'image.pdf')]
            for content,name in bad:
                r=client.post('/api/documents',files={'file':(name,content)})
                assert r.status_code==422,(name,r.status_code,r.text)
                if name=='image.pdf': assert 'OCR is unsupported' in r.text
            r=client.post('/api/documents',files={'file':('large.txt',b'a'*(10*1024*1024+1))});assert r.status_code==413
            doc=upload(data=b'Ordinary contract text.',name='../../outside.txt');assert doc['name']=='outside.txt'
            assert not (ROOT/'outside.txt').exists()
            assert client.post('/api/playbooks',headers={'Origin':'https://evil.invalid'},json={'yaml':'x'}).status_code==403
            assert client.get('/api/health',headers={'Host':'evil.invalid'}).status_code==400
            assert client.post('/api/playbooks',json={'yaml':'a: &a [*a]'}).status_code==422
        check('CR-05',invalid)
        def nda():
            path=ROOT/'corpus/contracts/01-acceptable-mutual-nda.txt'
            page.locator('.upload-box input').set_input_files(path)
            expect(page.get_by_role('heading',level=2,name=path.name)).to_be_visible()
            page.get_by_label('Review against').select_option('nda');page.get_by_role('button',name='Run review',exact=True).click()
            expect(page.locator('.issue-card')).to_have_count(0)
            expect(page.get_by_text('Review complete.',exact=False)).to_be_visible()
            page.get_by_role('tab',name='Rule checks',exact=False).click()
            expect(page.locator('.outcome.compliant')).to_have_count(6)
        check('CR-06',nda)
        def renewal():
            rev=state['review'];assert {i['rule_id'] for i in rev['issues']}=={'automatic-renewal','audit-rights'};validate(state['text'],rev)
            page.locator('.doc-row').filter(has_text=state['text']['name']).click()
            page.get_by_role('tab',name='Findings',exact=False).click()
            expect(page.locator('.issue-card')).to_have_count(2)
        check('CR-07',renewal)
        def liability():
            rev=review(upload(ROOT/'corpus/contracts/05-fee-based-liability-cap.txt'))
            cap=next(c for c in rev['checks'] if c['rule_id']=='liability-cap')
            assert cap['outcome']=='COMPLIANT' and cap['properties']['explicit_fee_cap'] and cap['properties']['period_months']==12
            exceptions=next(c for c in rev['checks'] if c['rule_id']=='liability-exceptions');assert exceptions['outcome']=='COMPLIANT'
        check('CR-08',liability)
        def absence():
            doc=upload(ROOT/'corpus/contracts/09-data-processing-addendum-missing.txt');rev=review(doc);i=byrule(rev,'data-protection');assert i['absence'] and i['observed_language'] is None and i['source_anchor'] is None;validate(doc,rev)
        check('CR-09',absence)
        def ip():
            doc=upload(ROOT/'corpus/contracts/07-broad-ip-assignment.txt');rev=review(doc);assert 'pre-existing' in byrule(rev,'ip-ownership')['observed_language'];validate(doc,rev)
        check('CR-10',ip)
        def edit():
            page.get_by_role('tab',name='Findings',exact=False).click()
            page.get_by_label('Review against').select_option('commercial')
            page.get_by_role('button',name='Duplicate',exact=True).click()
            page.get_by_role('button',name='Advanced YAML',exact=True).click()
            editor=page.get_by_label('Playbook YAML');expect(editor).to_be_visible()
            text=editor.input_value();assert 'severity: medium' in text
            editor.fill(text.replace('severity: medium','severity: low').replace('fallback_language: null', 'fallback_language: Either party may opt out by giving 45 days written notice.'))
            page.get_by_role('button',name='Save settings',exact=True).click()
            expect(page.get_by_role('dialog')).not_to_be_visible()
            page.get_by_role('button',name='Run review',exact=True).click()
            expect(page.locator('.issue-card .severity-label.low')).to_have_count(2)
            before=state['review'];after=client.get(f"/api/documents/{state['text']['id']}/reviews").json()[0]
            assert all(i['severity']=='low' for i in after['issues']) and all(i['severity']=='medium' for i in before['issues'])
            assert all(i['playbook_guidance']['fallback_language']=='Either party may opt out by giving 45 days written notice.' for i in after['issues'])
            assert all(i['playbook_guidance']['fallback_language'] is None for i in getrev(before['id'])['issues'])
            page.locator('.playbook-guidance summary').click()
            expect(page.locator('.playbook-guidance')).to_contain_text('Either party may opt out by giving 45 days written notice.')
            # Changing the policy later must not rewrite a completed review's guidance.
            book=client.get('/api/playbooks/'+after['playbook_id']).json()
            changed=book['yaml'].replace('45 days written notice', '75 days written notice')
            response=client.put('/api/playbooks/'+after['playbook_id'],json={'yaml':changed});response.raise_for_status()
            assert getrev(after['id'])['issues']==after['issues']
            page.reload()
            page.locator('.doc-row').filter(has_text=state['text']['name']).click()
            expect(page.locator('.issue-card .severity-label.low')).to_have_count(2)
            page.locator('.playbook-guidance summary').click()
            expect(page.locator('.playbook-guidance')).to_contain_text('45 days written notice')
            expect(page.locator('.playbook-guidance')).not_to_contain_text('75 days written notice')
            state['review']=after
            # Actual UI import/new validation and export.
            page.locator('.import-button input').set_input_files(ROOT/'playbooks/nda.yaml')
            expect(editor).to_be_visible();editor.fill(editor.input_value().replace('Mutual NDA','Imported NDA'))
            page.get_by_role('button',name='Save playbook',exact=True).click();expect(page.get_by_role('dialog')).not_to_be_visible()
            with page.expect_download() as download: page.get_by_role('button',name='Export YAML',exact=True).click()
            download.value.save_as(OUT/'export-playbook.yaml');assert 'Imported NDA' in (OUT/'export-playbook.yaml').read_text()
        check('CR-11',edit)
        def decisions():
            page.locator('.issue-card').first.click()
            page.get_by_role('button',name='Confirm finding',exact=True).click();expect(page.locator('.issue-card').first.locator('.issue-status')).to_have_text('Confirmed')
            page.get_by_label('Reviewer notes').fill('Discuss reciprocal audit scope with counsel.')
            page.get_by_role('button',name='Save note',exact=True).click()
            expect(page.get_by_role('status').filter(has_text='Reviewer update saved locally.')).to_be_visible()
            page.locator('.issue-card').nth(1).click();page.get_by_role('button',name='Dismiss finding',exact=True).click()
            expect(page.locator('.issue-card').nth(1).locator('.issue-status')).to_have_text('Dismissed')
            page.get_by_role('button',name='Clear decision',exact=True).click()
            expect(page.locator('.issue-card').nth(1).locator('.issue-status')).to_have_text('Needs decision')
            page.get_by_role('button',name='Dismiss finding',exact=True).click();expect(page.locator('.issue-card').nth(1).locator('.issue-status')).to_have_text('Dismissed')
            start();page.reload();page.locator('.doc-row').filter(has_text=state['text']['name']).click()
            expect(page.locator('.issue-card')).to_have_count(2)
            rev=getrev(state['review']['id']);assert {i['status'] for i in rev['issues']}=={'accepted','rejected'}
            assert any('Discuss reciprocal' in i['notes'] for i in rev['issues'])
        check('CR-12',decisions)
        def source():
            page.locator('.issue-card').first.click()
            quote=page.locator('.finding-detail blockquote').inner_text()
            assert quote in state['text']['text']
            page.get_by_role('button',name='View in source',exact=False).click()
            expect(page.locator('.source-highlight p')).to_have_text(quote)
            page.screenshot(path=str(OUT/'source-navigation.png'),full_page=True)
            page.get_by_role('tab',name='Findings',exact=False).click()
        check('CR-13',source)
        def exports():
            rev=getrev(state['review']['id'])
            for fmt,ext in [('json','json'),('markdown','md'),('docx','docx')]:
                response=client.get(f"/api/reviews/{rev['id']}/export",params={'format':fmt});response.raise_for_status();(OUT/f'review.{ext}').write_bytes(response.content)
                if fmt=='json':
                    data=response.json();assert data['issues']==rev['issues'];assert data['author']=='Ethan'
                    required={'id','rule_id','severity','category','observed_language','source_anchor','explanation','preferred_position','negotiating_response','status','notes','confidence','absence'}
                    for issue in data['issues']: assert required.issubset(issue)
                else:
                    if fmt=='docx':
                        doc=Document(BytesIO(response.content));content='\n'.join(p.text for p in doc.paragraphs);assert doc.core_properties.author==doc.core_properties.last_modified_by=='Ethan'
                    else: content=response.text
                    for label in ['Issue ID','Rule ID','Severity','Clause category','Observed language','Source anchor','Explanation','Preferred contract position','Suggested negotiating response','Reviewer status','Reviewer notes','Confidence or uncertainty']: assert label in content,label
                    for issue in rev['issues']: assert issue['id'] in content and issue['rule_id'] in content
                    assert 'Saved playbook alternatives and fallback' in content and '45 days written notice' in content and '75 days written notice' not in content
            # Click all browser download controls as well.
            for name in ['Markdown','DOCX','JSON']:
                with page.expect_download() as download: page.get_by_label('Export review').select_option(name.lower())
                assert download.value.suggested_filename
        check('CR-14',exports)
        def deletion():
            doc=state['text'];rid=state['review']['id']
            page.once('dialog',lambda d:d.accept());page.get_by_role('button',name='Delete',exact=True).click()
            expect(page.locator('.doc-row').filter(has_text=doc['name'])).to_have_count(0)
            assert client.get('/api/documents/'+doc['id']).status_code==404 and client.get('/api/reviews/'+rid).status_code==404
            assert not (DATA/'sources'/doc['id']).exists()
            with sqlite3.connect(DATA/'workbench.sqlite3') as c: assert c.execute('SELECT count(*) FROM reviews WHERE document_id=?',(doc['id'],)).fetchone()[0]==0
        check('CR-15',deletion)
        def corpus():
            r=evaluate(BASE,OUT/'corpus-results.json');assert r['passed']==r['total']==12, json.dumps(r)
        check('CR-16',corpus)
        def privacy():
            assert not client.get('/api/health').json()['ai_configured']
            assert NETWORK.read_text()=='',NETWORK.read_text()
            assert not external,external
            doc=upload(data=b'Provider may amend unilaterally without consent.');rev=review(doc)
            r=client.post(f"/api/reviews/{rev['id']}/ai",json={'consent':False,'model':'openai/gpt-4o-mini','provider':'OpenAI'});assert r.status_code==403
            r=client.post(f"/api/reviews/{rev['id']}/ai",json={'consent':True,'model':'openai/gpt-4o-mini','provider':'OpenAI'});assert r.status_code==502 and 'BLOCKED_EXTERNAL' in r.text
            assert NETWORK.read_text()==''
            return 'No backend outbound socket connections or external browser requests during offline ingestion, review, corpus and exports; consent and missing-key gates verified.'
        check('CR-17',privacy)
        def invalid_ai():
            start(key='invalid-acceptance-credential')
            doc=upload(data=b'Each party protects confidential information.');rev=review(doc)
            r=client.post(f"/api/reviews/{rev['id']}/ai",json={'consent':True,'model':'openai/gpt-4o-mini','provider':'OpenAI'})
            assert r.status_code==502 and ('HTTP 401' in r.text or 'HTTP 403' in r.text),r.text
            assert getrev(rev['id'])['issues']==rev['issues'] and getrev(rev['id'])['ai_status']=='ERROR'
            assert client.get('/api/documents/'+doc['id']).status_code==200
            assert NETWORK.read_text().strip(),'Must be an actual provider request'
            return 'Real OpenRouter request with deliberately invalid credential rejected; source and deterministic review preserved.'
        if os.environ.get('CRW_ALLOW_PROVIDER_REJECTION_TEST') == '1':
            check('CR-19',invalid_ai)
        else:
            results.append({'id':'CR-19','status':'BLOCKED_EXTERNAL','detail':'Provider rejection test requires explicit authorization and CRW_ALLOW_PROVIDER_REJECTION_TEST=1. No provider request made.'})
        def fresh():
            freshdata=RUN/'fresh';start(data=freshdata)
            assert client.get('/api/documents').json()==[]
            page.goto(BASE);unseen=RUN/'unseen.txt';unseen.write_text('ORBIT SERVICES AGREEMENT\n\nProvider may amend the service agreement unilaterally without customer consent.\n\nEach party retains its pre-existing intellectual property.')
            page.locator('.upload-box input').set_input_files(unseen);expect(page.get_by_role('heading',level=2,name='unseen.txt')).to_be_visible()
            page.get_by_label('Review against').select_option('commercial');page.get_by_role('button',name='Run review',exact=True).click();expect(page.locator('.issue-card').first).to_be_visible()
            page.locator('.issue-card').filter(has_text='unilateral-amendment').click();page.get_by_role('button',name='Confirm finding',exact=True).click()
            expect(page.locator('.issue-card').filter(has_text='unilateral-amendment').locator('.issue-status')).to_have_text('Confirmed')
            with page.expect_download() as download: page.get_by_label('Export review').select_option('json')
            download.value.save_as(OUT/'fresh-report.json');report=json.loads((OUT/'fresh-report.json').read_text());assert byrule(report,'unilateral-amendment')['status']=='accepted'
            page.screenshot(path=str(OUT/'workbench.png'),full_page=True)
            page.set_viewport_size({'width':390,'height':844});page.screenshot(path=str(OUT/'mobile.png'),full_page=True)
            assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
        check('CR-20',fresh)
        def additional():
            # Ambiguity and conflicting provisions must not yield a clean cap check.
            base=(ROOT/'corpus/contracts/05-fee-based-liability-cap.txt').read_text()
            doc=upload(data=(base+'\n\nAggregate liability is unlimited.').encode())
            rev=review(doc);assert byrule(rev,'liability-cap')['confidence']=='OBJECTIVE';validate(doc,rev)
            malicious=upload(data=b'<script>window.injected=true</script>\nProvider may amend unilaterally without consent.',name='safe.txt')
            page.reload();page.locator('.doc-row').filter(has_text='safe.txt').click()
            page.get_by_role('tab',name='Extracted source',exact=False).click();expect(page.locator('.source-block').first).to_contain_text('<script>')
            assert page.evaluate('window.injected') is None
            # Age one record in the real SQLite DB then restart with retention enabled.
            with sqlite3.connect(RUN/'fresh'/'workbench.sqlite3') as c:
                c.execute('UPDATE documents SET created=? WHERE id=?',(time.time()-3*86400,doc['id']))
            os.environ['CRW_RETENTION_DAYS']='1'
            try: start(data=RUN/'fresh')
            finally: os.environ.pop('CRW_RETENTION_DAYS',None)
            assert client.get('/api/documents/'+doc['id']).status_code==404
            assert client.get('/api/reviews/'+rev['id']).status_code==404
            assert not (RUN/'fresh'/'sources'/doc['id']).exists()
            return 'Real API conflict detection, literal script rendering, and retention cascade verified.'
        check('CR-EXTRA',additional)
        key=os.environ.get('OPENROUTER_API_KEY');model=os.environ.get('OPENROUTER_MODEL');provider=os.environ.get('OPENROUTER_PROVIDER')
        if os.environ.get('CRW_ALLOW_LIVE_AI_TEST') == '1' and key and model and provider:
            def live_ai():
                start(key=key);doc=upload(ROOT/'corpus/contracts/12-ambiguous-liability-cap.txt');rev=review(doc)
                r=client.post(f"/api/reviews/{rev['id']}/ai",json={'consent':True,'model':model,'provider':provider,'max_cost_usd':0.10});r.raise_for_status();validate(doc,r.json());assert r.json()['ai_status']=='VERIFIED_CITATIONS_REQUIRES_REVIEW'
            check('CR-18',live_ai)
        else:
            results.append({'id':'CR-18','status':'BLOCKED_EXTERNAL','detail':'Live AI requires explicit authorization, CRW_ALLOW_LIVE_AI_TEST=1 and provider configuration. No provider request or fake response used.'})
        browser.close()
finally:
    stop();client.close();REGISTRY.write_text('[]')
    (OUT/'e2e-results.json').write_text(json.dumps({'timestamp':datetime.now(timezone.utc).isoformat(),'run':str(RUN.relative_to(ROOT)),'results':sorted(results,key=lambda r:r['id'])},indent=2)+'\n')
raise SystemExit(1 if any(r['status']=='FAIL' for r in results) else 0)
